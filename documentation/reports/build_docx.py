from __future__ import annotations

import re
import sys
from pathlib import Path

from docx import Document
from docx.document import Document as DocumentObject
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.shared import Inches, Pt, RGBColor
from docx.text.paragraph import Paragraph


INLINE_MARKDOWN = re.compile(r"(\*\*.*?\*\*|`[^`]+`|\[[^\]]+\]\([^)]+\))")
TABLE_SEPARATOR = re.compile(r"^\|?\s*:?-{3,}:?\s*(\|\s*:?-{3,}:?\s*)+\|?$")


def _add_inline_content(paragraph: Paragraph, content: str) -> None:
    cursor = 0
    for match in INLINE_MARKDOWN.finditer(content):
        if match.start() > cursor:
            paragraph.add_run(content[cursor : match.start()])

        token = match.group()
        if token.startswith("**"):
            run = paragraph.add_run(token[2:-2])
            run.bold = True
        elif token.startswith("`"):
            run = paragraph.add_run(token[1:-1])
            run.font.name = "Consolas"
            run.font.color.rgb = RGBColor(46, 87, 117)
        else:
            link_match = re.fullmatch(r"\[([^\]]+)\]\(([^)]+)\)", token)
            if link_match is None:
                paragraph.add_run(token)
                cursor = match.end()
                continue
            label, url = link_match.groups()
            paragraph.add_run(f"{label} ({url})")
        cursor = match.end()

    if cursor < len(content):
        paragraph.add_run(content[cursor:])


def _table_cells(line: str) -> list[str]:
    return [cell.strip() for cell in line.strip().strip("|").split("|")]


def _add_table(document: DocumentObject, rows: list[list[str]]) -> None:
    column_count = len(rows[0])
    table = document.add_table(rows=0, cols=column_count)
    table.style = "Light Shading Accent 1"

    for row_index, row_values in enumerate(rows):
        cells = table.add_row().cells
        for column_index, cell_text in enumerate(row_values):
            cell = cells[column_index]
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            paragraph = cell.paragraphs[0]
            _add_inline_content(paragraph, cell_text)
            for run in paragraph.runs:
                run.font.size = Pt(8.5)
                if row_index == 0:
                    run.bold = True


def build_docx(markdown_path: Path, output_path: Path) -> None:
    document = Document()
    section = document.sections[0]
    section.top_margin = Inches(0.65)
    section.bottom_margin = Inches(0.65)
    section.left_margin = Inches(0.7)
    section.right_margin = Inches(0.7)

    normal_style = document.styles["Normal"]
    normal_style.font.name = "Aptos"
    normal_style.font.size = Pt(10)
    normal_style.paragraph_format.space_after = Pt(5)

    title_properties = document.core_properties
    title_properties.title = "Relatório Técnico — Armazenamento em Nuvem na AWS — V1"
    title_properties.subject = "Atividade 3: S3, RDS e DynamoDB"
    title_properties.author = "Grupo Ctrl Alt Del — Shifu"

    header = section.header.paragraphs[0]
    header.text = "SHIFU  •  COMPUTAÇÃO EM NUVEM II  •  ATIVIDADE V1"
    header.runs[0].font.size = Pt(8)
    header.runs[0].font.color.rgb = RGBColor(82, 101, 120)

    footer = section.footer.paragraphs[0]
    footer.text = "Relatório acadêmico — revisar resultados e evidências após execução na AWS"
    footer.runs[0].font.size = Pt(8)
    footer.runs[0].font.color.rgb = RGBColor(82, 101, 120)

    lines = markdown_path.read_text(encoding="utf-8").splitlines()
    index = 0
    inside_code = False
    code_lines: list[str] = []

    while index < len(lines):
        line = lines[index]

        if line.startswith("```"):
            if inside_code:
                paragraph = document.add_paragraph()
                paragraph.paragraph_format.left_indent = Inches(0.2)
                paragraph.paragraph_format.space_after = Pt(7)
                run = paragraph.add_run("\n".join(code_lines))
                run.font.name = "Consolas"
                run.font.size = Pt(8)
                run.font.color.rgb = RGBColor(55, 65, 81)
                code_lines = []
                inside_code = False
            else:
                inside_code = True
            index += 1
            continue

        if inside_code:
            code_lines.append(line)
            index += 1
            continue

        if line.startswith("|") and index + 1 < len(lines) and TABLE_SEPARATOR.match(lines[index + 1]):
            rows = [_table_cells(line)]
            index += 2
            while index < len(lines) and lines[index].startswith("|"):
                rows.append(_table_cells(lines[index]))
                index += 1
            _add_table(document, rows)
            continue

        if line.startswith("#"):
            heading_marker, heading_text = line.split(" ", 1)
            level = len(heading_marker)
            document.add_heading(heading_text, level=0 if level == 1 else min(level - 1, 3))
        elif line.startswith("> "):
            paragraph = document.add_paragraph(style="Intense Quote")
            _add_inline_content(paragraph, line[2:])
        elif line.startswith("- "):
            paragraph = document.add_paragraph(style="List Bullet")
            _add_inline_content(paragraph, line[2:])
        elif re.match(r"^\d+\.\s", line):
            paragraph = document.add_paragraph(style="List Number")
            _add_inline_content(paragraph, re.sub(r"^\d+\.\s", "", line))
        elif line.strip():
            paragraph = document.add_paragraph()
            _add_inline_content(paragraph, line.replace("  ", " "))

        index += 1

    output_path.parent.mkdir(parents=True, exist_ok=True)
    document.save(output_path)


if __name__ == "__main__":
    report_path = Path(__file__).with_name("armazenamento-v1.md")
    destination = Path(sys.argv[1]) if len(sys.argv) > 1 else report_path.with_suffix(".docx")
    build_docx(report_path, destination)
    print(f"Relatorio DOCX gerado: {destination.resolve()}")
